"""One customer's consumption as an Excel workbook: values and native charts,
no formulas (operator, 2026-10-16). Modelled on the operator's own workbook --
a Dashboard of KPIs and charts over sheets holding the data each chart reads.

`openpyxl` is provided to plugin backends by the host (framework API 9).
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY = '#,##0;[Red]-#,##0'
PERCENT = '0.0"%"'
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16)
LABEL_FONT = Font(bold=True, color="595959", size=9)
KPI_FONT = Font(bold=True, size=14)
TOP_NAMES = 15


def _table(sheet, top_row: int, headers: list[str], rows: list[list], formats: dict[int, str] | None = None) -> int:
    """Writes a header row and rows from `top_row`; returns the last row written."""
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(top_row, column, header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for offset, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(top_row + offset, column, value)
            if formats and column in formats and isinstance(value, (int, float)):
                cell.number_format = formats[column]
    for column in range(1, len(headers) + 1):
        width = max([len(str(headers[column - 1]))] + [len(str(row[column - 1])) for row in rows
                                                       if column - 1 < len(row) and row[column - 1] is not None])
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 10), 48)
    sheet.freeze_panes = sheet.cell(top_row + 1, 2)
    return top_row + len(rows)


def _bar(title: str, sheet, first_row: int, last_row: int, value_columns: list[int], label_column: int = 1,
         stacked: bool = False, percent: bool = False, height: float = 7.5, width: float = 16) -> BarChart:
    chart = BarChart()
    chart.type, chart.title, chart.height, chart.width = "col", title, height, width
    if stacked:
        chart.grouping, chart.overlap = "stacked", 100
    for column in value_columns:
        chart.add_data(Reference(sheet, min_col=column, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=label_column, min_row=first_row, max_row=last_row))
    chart.y_axis.numFmt = '0"%"' if percent else '#,##0'
    chart.y_axis.majorGridlines = None if percent else chart.y_axis.majorGridlines
    chart.legend.position = "b"
    return chart


def _line(title: str, sheet, first_row: int, last_row: int, value_columns: list[int], label_column: int = 1,
          height: float = 7.5, width: float = 16) -> LineChart:
    chart = LineChart()
    chart.title, chart.height, chart.width = title, height, width
    for column in value_columns:
        chart.add_data(Reference(sheet, min_col=column, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=label_column, min_row=first_row, max_row=last_row))
    chart.y_axis.numFmt = '#,##0'
    chart.legend.position = "b"
    return chart


def _top_names(merged: dict, kind_key: str) -> list[str]:
    totals: dict[str, float] = {}
    for data in merged.values():
        for name, cost in ((data.get("actual") or {}).get(kind_key) or {}).items():
            totals[name] = totals.get(name, 0.0) + cost
    return [name for name, _ in sorted(totals.items(), key=lambda item: -item[1])]


def _breakdown_sheet(workbook: Workbook, title: str, merged: dict, kind_key: str, limit: int | None = TOP_NAMES):
    sheet = workbook.create_sheet(title)
    names = _top_names(merged, kind_key)
    shown = names[:limit] if limit else names
    rows = []
    for month in sorted(merged):
        spend = (merged[month].get("actual") or {}).get(kind_key) or {}
        row = [month] + [round(spend.get(name, 0.0), 2) for name in shown]
        if limit and len(names) > limit:
            row.append(round(sum(cost for name, cost in spend.items() if name not in shown), 2))
        rows.append(row)
    headers = ["Month"] + shown + (["Other"] if limit and len(names) > limit else [])
    last = _table(sheet, 1, headers, rows, {column: MONEY for column in range(2, len(headers) + 1)})
    return sheet, last, len(headers)


def build(customer: dict, analysis: dict, merged: dict, today: date) -> bytes:
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    currency = analysis.get("currency") or "USD"
    board = analysis["dashboard"]
    kpis = board["kpis"]
    measure = analysis["measure"]

    # -- data sheets ------------------------------------------------------
    monthly = workbook.create_sheet("Monthly")
    month_rows = [[row["month"], row["usage"], row["marketplace"], row["credits"], row["actual"], row["amortized"],
                   row["mom_pct"], row["yoy_pct"], "partial" if row["partial"] else ""] for row in board["monthly"]]
    monthly_last = _table(monthly, 1, ["Month", "Azure usage", "Marketplace", "Credits", "Total", "Amortized",
                                       "MoM %", "YoY %", ""], month_rows,
                          {2: MONEY, 3: MONEY, 4: MONEY, 5: MONEY, 6: MONEY, 7: PERCENT, 8: PERCENT})

    mix = workbook.create_sheet("Mix")
    lines = sorted({name for row in board["mix"] for name in row["lines"]},
                   key=lambda name: (name == "Azure (other)", name))
    mix_last = _table(mix, 1, ["Month"] + lines,
                      [[row["month"]] + [row["lines"].get(name, 0.0) for name in lines] for row in board["mix"]],
                      {column: MONEY for column in range(2, len(lines) + 2)})

    years = workbook.create_sheet("Years")
    years_last = _table(years, 1, ["Year", "Total", "YTD (same months)"],
                        [[row["year"], row["total"], next((y["total"] for y in board["ytd_by_year"]
                                                           if y["year"] == row["year"]), None)]
                         for row in board["by_year"]], {2: MONEY, 3: MONEY})
    comparison = board["this_vs_last_year"]
    _table(years, years_last + 3, ["Month", str(comparison["last_year"]), str(comparison["this_year"]),
                                   f"MoM % {comparison['last_year']}", f"MoM % {comparison['this_year']}"],
           [[name, comparison["last"][index], comparison["this"][index], comparison["last_mom"][index],
             comparison["this_mom"][index]]
            for index, name in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct",
                                          "Nov", "Dec"])],
           {2: MONEY, 3: MONEY, 4: PERCENT, 5: PERCENT})
    comparison_first = years_last + 4

    _breakdown_sheet(workbook, "Services", merged, "service")
    _breakdown_sheet(workbook, "Subscriptions", merged, "subscription")
    enrollments, enrollments_last, enrollments_columns = _breakdown_sheet(workbook, "Enrollments", merged,
                                                                          "account", limit=None)
    quarter_rows = [[row["quarter"]] + [row["by_account"].get(name, 0.0) for name in board["accounts"]]
                    for row in board["quarterly"]]
    quarter_first = enrollments_last + 3
    _table(enrollments, quarter_first, ["Quarter"] + board["accounts"], quarter_rows,
           {column: MONEY for column in range(2, len(board["accounts"]) + 2)})

    ai = workbook.create_sheet("GPU & AI")
    _table(ai, 1, ["Month", "Line", "Bought as", "Cost"],
           [[row["month"], line["category"], line["pricing"], line["cost"]]
            for row in analysis["gpu_and_ai"] for line in row["lines"]], {4: MONEY})

    forecast_sheet = workbook.create_sheet("Forecast")
    automatic = analysis.get("automatic_forecast")
    if automatic:
        summary = [["Method", automatic["method"]], ["Months of history", automatic["months_of_history"]],
                   ["Fit (R²)", automatic["r_squared"]], ["Backtest error (MAPE)", automatic["mape"]],
                   ["Next 12 months", automatic["next_12_months"]],
                   ["Next 12 months + pipeline", automatic["next_12_months_with_pipeline"]],
                   [f"{today.year} total (actual + forecast)", automatic["this_year_total"]],
                   [f"{automatic['next_year']} forecast", automatic["next_year_total"]],
                   [f"{automatic['next_year']} forecast + pipeline", automatic["next_year_with_pipeline"]]]
        for offset, (label, value) in enumerate(summary, start=1):
            forecast_sheet.cell(offset, 1, label).font = LABEL_FONT
            cell = forecast_sheet.cell(offset, 2, value)
            if isinstance(value, float) and value > 10:
                cell.number_format = MONEY
        forecast_first = len(summary) + 3
        forecast_last = _table(forecast_sheet, forecast_first, ["Month", "Actual", "Model", "Forecast", "Pipeline"],
                               [[row["month"], row["actual"], row["model"], row["forecast"], row["pipeline"] or None]
                                for row in automatic["series"]], {2: MONEY, 3: MONEY, 4: MONEY, 5: MONEY})
        forecast_sheet.add_chart(_line("Actual vs model vs forecast", forecast_sheet, forecast_first + 1,
                                       forecast_last, [2, 3, 4], width=24), "H2")
    else:
        forecast_sheet["A1"] = "Not enough billed months for an automatic forecast (six are needed)."

    scenario_sheet = workbook.create_sheet("Scenarios")
    scenarios = analysis.get("scenarios")
    if scenarios:
        settings = scenarios["settings"]
        scenario_sheet["A1"] = (f"From {scenarios['base_month']} ({scenarios['base_total']:,.0f}); monthly growth "
                                + ", ".join(f"{name} {rate * 100:.1f}%" for name, rate in settings["growth"].items())
                                + f"; seasonality {'on' if settings['seasonality'] else 'off'}; active {settings['active']}"
                                + (f"; uplift {settings['uplift_usd']:,.0f}/month from {settings['uplift_from']}"
                                   if settings.get("uplift_from") else ""))
        scenario_last = _table(scenario_sheet, 3, ["Month", "Base", "Upside", "Downside", "Pipeline",
                                                   f"{settings['active']} + pipeline"],
                               [[row["month"], row["Base"], row["Upside"], row["Downside"], row["pipeline"],
                                 row["active_with_pipeline"]] for row in scenarios["rows"]],
                               {column: MONEY for column in range(2, 7)})
        scenario_sheet.add_chart(_line("Scenarios", scenario_sheet, 4, scenario_last, [2, 3, 4, 6], width=24), "H3")

    pipeline_sheet = workbook.create_sheet("Pipeline")
    pipeline = analysis.get("pipeline") or {}
    opportunities = pipeline.get("opportunities") or []
    if opportunities:
        pipeline_last = _table(pipeline_sheet, 1, ["Opportunity", "Under", "Status", "Stage", "Partner",
                                                   "Billing profile", "Starts", "Expected monthly"],
                               [[o.get("title"), o.get("affiliate") or o.get("customer"), o.get("status"),
                                 o.get("stage"), o.get("partner"), o.get("billing_profile"),
                                 o.get("consumption_start"), o.get("expected_monthly_usd")] for o in opportunities],
                               {8: MONEY})
        plan_months = sorted({row["month"] for o in opportunities for row in o.get("plan") or []})
        if plan_months:
            _table(pipeline_sheet, pipeline_last + 3, ["Opportunity"] + plan_months,
                   [[o.get("title")] + [next((row["usd"] for row in o.get("plan") or [] if row["month"] == month), None)
                                        for month in plan_months] for o in opportunities],
                   {column: MONEY for column in range(2, len(plan_months) + 2)})
    else:
        pipeline_sheet["A1"] = pipeline.get("reason") or "No opportunities for this customer."

    commitment = analysis.get("commitment")
    if commitment:
        commitment_sheet = workbook.create_sheet("Commitment")
        facts = [["Signed", commitment["amount"]], ["Type", commitment["type"]],
                 ["Term", f"{commitment['start']} to {commitment['end']}"],
                 ["Counts Marketplace", "yes" if commitment.get("include_marketplace", True) else "no"],
                 ["Subtracts credits", "yes" if commitment.get("subtract_credits", True) else "no"],
                 ["Consumed", commitment["consumed"]], ["Consumed %", commitment["consumed_pct"]],
                 ["Term elapsed %", commitment["elapsed_pct"]], ["Remaining", commitment["remaining"]],
                 ["Against an even pace", commitment["pace_gap"]], ["Run rate (3 closed months)", commitment["run_rate"]],
                 ["At term end, at run rate", commitment["projected_at_end"]],
                 ["At term end, with pipeline", commitment["projected_with_pipeline"]],
                 ["Projected shortfall", commitment["projected_shortfall"]],
                 ["Fully consumed by", commitment["projected_exhaustion"] or "not within the term"],
                 ["Verdict", commitment["verdict"]]]
        for offset, (label, value) in enumerate(facts, start=1):
            commitment_sheet.cell(offset, 1, label).font = LABEL_FONT
            cell = commitment_sheet.cell(offset, 2, value)
            if isinstance(value, float) and abs(value) > 100:
                cell.number_format = MONEY
        cumulative_first = len(facts) + 3
        cumulative_last = _table(commitment_sheet, cumulative_first, ["Month", "Consumed to date", "Signed"],
                                 [[row["month"], row["consumed"], commitment["amount"]] for row in commitment["cumulative"]],
                                 {2: MONEY, 3: MONEY})
        commitment_sheet.add_chart(_line("Consumed against the signed value", commitment_sheet, cumulative_first + 1,
                                         cumulative_last, [2, 3]), "E2")

    if board["credit_details"]:
        credits = workbook.create_sheet("Credits")
        _table(credits, 1, ["Month", "Credit", "Value"],
               [[row["month"], row["name"], row["value"]] for row in board["credit_details"]], {3: MONEY})

    # -- dashboard ----------------------------------------------------------
    dashboard["A1"] = f"Azure consumption -- {customer['name']}"
    dashboard["A1"].font = TITLE_FONT
    dashboard["A2"] = (f"{len(customer['accounts'])} enrollment(s) · {kpis['months_of_history']} months from "
                       f"{kpis['first_month']} · through {kpis['latest_month']} (closed) · {currency} · "
                       f"Marketplace {'included' if measure['marketplace'] else 'excluded'}, credits "
                       f"{'subtracted' if measure['credits'] else 'not subtracted'} · exported {today.isoformat()}")
    tiles = [("Total, all synced months", kpis["all_time"], MONEY),
             (f"YTD through {kpis['ytd_through']}", kpis["ytd"], MONEY),
             ("YTD against last year", kpis["ytd_growth_pct"], PERCENT),
             (f"Latest month ({kpis['latest_month']})", kpis["latest_total"], MONEY),
             ("Latest month MoM", kpis["latest_mom_pct"], PERCENT),
             ("Latest month YoY", kpis["latest_yoy_pct"], PERCENT),
             ("Next 12 months (auto forecast)", kpis["next_12_months"], MONEY),
             ("Next year (auto forecast)", kpis["next_year_forecast"], MONEY),
             ("Burn rate, 3-month average", kpis["burn_rate_3m"], MONEY),
             ("Largest enrollment's share", kpis["largest_enrollment_share_pct"], PERCENT),
             ("Credits applied, all synced months", kpis["credits_all_time"], MONEY),
             ("Marketplace, all synced months", kpis["marketplace_all_time"], MONEY)]
    if commitment:
        tiles += [(f"Signed ({commitment['type']})", commitment["amount"], MONEY),
                  ("Consumed against it", commitment["consumed_pct"], PERCENT),
                  ("Remaining", commitment["remaining"], MONEY),
                  ("Runway at current burn (months)", kpis.get("runway_months"), "0.0"),
                  ("Projected exhaustion", commitment["projected_exhaustion"] or "not within term", None)]
    for index, (label, value, number_format) in enumerate(tiles):
        row, column = 4 + (index // 6) * 3, 1 + (index % 6) * 2
        dashboard.cell(row, column, label).font = LABEL_FONT
        cell = dashboard.cell(row + 1, column, value if value is not None else "—")
        cell.font = KPI_FONT
        if number_format and isinstance(value, (int, float)):
            cell.number_format = number_format
    for column in range(1, 13):
        dashboard.column_dimensions[get_column_letter(column)].width = 16

    anchor_row = 6 + ((len(tiles) - 1) // 6 + 1) * 3
    charts = [
        _bar("Total consumption by month", monthly, 2, monthly_last, [5]),
        _bar("Mix by month", mix, 2, mix_last, list(range(2, len(lines) + 2)), stacked=True),
        _bar("By year", years, 2, years_last, [2]),
        _bar("Year to date, same months", years, 2, years_last, [3]),
        _bar("Month over month %", monthly, 2, monthly_last, [7], percent=True),
        _bar("Year over year %", monthly, 2, monthly_last, [8], percent=True),
        _bar(f"{comparison['last_year']} against {comparison['this_year']}, by month", years, comparison_first + 1,
             comparison_first + 12, [2, 3]),
        _bar(f"MoM growth % {comparison['last_year']} against {comparison['this_year']}", years,
             comparison_first + 1, comparison_first + 12, [4, 5], percent=True),
        _bar("By enrollment, by month", enrollments, 2, enrollments_last, list(range(2, enrollments_columns + 1)),
             stacked=True),
    ]
    if quarter_rows:
        charts.append(_bar("By enrollment, by quarter", enrollments, quarter_first + 1, quarter_first + len(quarter_rows),
                           list(range(2, len(board["accounts"]) + 2))))
    for index, chart in enumerate(charts):
        row = anchor_row + (index // 2) * 16
        column = "A" if index % 2 == 0 else "G"
        dashboard.add_chart(chart, f"{column}{row}")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
