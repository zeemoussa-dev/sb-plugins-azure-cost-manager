"""The cross-customer reports' workbooks.

Worth pinning: each opens, carries the template's navy summary tab with its
sheets in order, and puts money and percentages in the template's formats.
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import load_workbook

from backend import export_reports
from backend.export_xlsx import MONEY, NAVY, PERCENT

TODAY = date(2026, 10, 21)


def opened(content: bytes):
    return load_workbook(io.BytesIO(content))


def test_growth_workbook():
    rows = [{"id": "fab", "name": "First Abu Dhabi Bank", "affiliates": 0, "status": "growing", "last_month": "2026-09",
             "last_total": 13_900_000.0, "recent_monthly": 7_400_000.0, "change_3m_pct": 133.7, "change_12m_pct": 133.0,
             "yoy_pct": 856.5, "signals": [{"kind": "growing", "text": "Up 133.7% over three months"}],
             "stopped": [{"service": "HDInsight", "before": 5000.0, "last": 10.0}],
             "started": [{"service": "Microsoft Fabric", "last": 9000.0}],
             "by_play": [{"play": "AI & ML", "recent": 100.0, "before": 50.0, "change": 50.0, "change_pct": 100.0}]},
            {"id": "x", "name": "Empty", "status": "no data", "signals": [], "last_month": None}]

    book = opened(export_reports.build_growth(rows, TODAY))

    assert book.sheetnames == ["Growth and Churn", "Stopped and New", "By Solution Play"]
    summary = book["Growth and Churn"]
    assert summary.sheet_properties.tabColor.rgb.endswith(NAVY)
    assert summary["A9"].value == "First Abu Dhabi Bank" and summary["D9"].number_format == MONEY
    assert summary["F9"].value == 1.337 and summary["F9"].number_format == PERCENT
    assert [book["Stopped and New"].cell(row, 3).value for row in (5, 6)] == ["HDInsight", "Microsoft Fabric"]


def test_whitespace_and_marketplace_whitespace_workbooks():
    whitespace = opened(export_reports.build_whitespace([
        {"id": "taqa", "name": "TAQA", "monthly": 2_000_000.0, "uses": 38, "entity_customer": None, "currency": "USD",
         "gaps": [{"service": "Azure Firewall", "play": "Security", "users": 20, "adoption_pct": 45,
                   "typical_monthly": 12_000.0}]}], TODAY))
    assert whitespace.sheetnames == ["Service Whitespace", "Gaps"]
    assert whitespace["Gaps"]["B5"].value == "Azure Firewall" and whitespace["Gaps"]["E5"].value == 0.45

    marketplace = opened(export_reports.build_marketplace_whitespace({
        "period": {"first": "2025-11", "last": "2026-10"}, "third_party": 800.0, "core42": 300.0,
        "core42_share_pct": 27.3,
        "categories": [{"category": "Security", "total": 500.0, "macc_eligible": 500.0, "customers": 2,
                        "solutions": [{"publisher": "Fortinet", "offer": "FortiGate", "total": 500.0, "customers": 2,
                                       "kind": "SaaS", "macc_eligible": True}]}],
        "customers": [{"name": "ADNOC", "third_party": 400.0, "macc_eligible": 400.0, "core42": 300.0,
                       "core42_share_pct": 42.9, "by_category": {"Security": 400.0}}]}, TODAY))
    assert marketplace.sheetnames == ["Marketplace Whitespace", "Leading Offers", "By Customer"]
    assert marketplace["By Customer"]["E4"].value == "Security" and marketplace["By Customer"]["E5"].value == 400.0


def test_hygiene_and_renewals_workbooks():
    hygiene = opened(export_reports.build_hygiene([
        {"id": "dge", "name": "DGE", "last_month": "2026-09", "subscriptions": 964, "subscriptions_12m": 1003, "tiny": 229,
         "top_share_pct": 3.7, "top3_share_pct": 9.0, "new": [{"subscription": "AI Lab", "last": 600.0}], "new_count": 12,
         "dormant": [{"subscription": "Old", "before": 800.0, "last": 20.0}], "dormant_count": 9, "tags_synced": True,
         "tag_coverage": {"Environment": 41.3, "Project": 10.5}}], TODAY))
    assert hygiene.sheetnames == ["Hygiene", "New and Dormant"]
    assert hygiene["Hygiene"]["I9"].value == 0.413 and hygiene["Hygiene"]["K9"].value is None

    renewals = opened(export_reports.build_renewals([
        {"customer": "ADNOC", "publisher": "Informatica", "offer": "IDMC", "plan": "Year 3", "subscription": "Prod",
         "core42": False, "cadence": "yearly", "cadence_months": 12, "cadence_assumed": False, "last_month": "2025-11",
         "last_amount": 373_301.0, "charges": 2, "next_month": "2026-11", "status": "upcoming"}], TODAY))
    assert renewals.sheetnames == ["Renewals", "Calendar"]
    assert renewals["Calendar"]["C5"].value == "ADNOC" and renewals["Calendar"]["I5"].number_format == MONEY
