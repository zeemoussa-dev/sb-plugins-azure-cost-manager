"""FinOps: coverage, steady on-demand candidates, purchases, growth and budget.

Worth pinning: coverage is read from amortized cost (under actual cost a
reservation's usage costs nothing); a candidate must be steady AND large; the
current month is never evidence; the budget is the MACC over its term unless
overridden; and the Customer view drops only the forecast scenarios.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timezone

import openpyxl

from backend import finops
from backend.consumption_service import ConsumptionService
from backend.consumption_store import ConsumptionStore
from backend.customers import AccountDirectory, CustomerBook

TODAY = date(2026, 10, 17)
CLOSED = [f"2026-{m:02d}" for m in range(1, 10)]
ACCOUNTS = [{"name": "1", "display_name": "Contoso", "agreement_type": "MicrosoftCustomerAgreement",
             "term_start": "2026-01-01", "term_end": "2028-12-31"}]


def month(meters: dict[str, float], amortized: dict[str, float] | None = None, total: float = 10_000.0,
          services: dict[str, float] | None = None) -> dict:
    services = services or {"Virtual Machines": total}
    spend = {"service": services, "subscription": {"Prod": total}}
    data = {"fetched_at": "2026-10-17T00:00:00+00:00", "currency": "USD", "actual": spend, "amortized": spend,
            "meters": meters}
    if amortized is not None:
        data["meters_amortized"] = {"by": amortized, "fetched_at": "2026-10-17T00:00:00+00:00"}
    return data


def entries_with(months: dict[str, dict]) -> dict[str, dict]:
    return {"1": {"months": months}}


def test_coverage_is_read_from_amortized_cost_by_family():
    amortized = {"Dsv5 Series|OnDemand": 3_000, "Dsv5 Series|Reservation": 1_000,
                 "NDsr H100 v5 Series|SavingsPlan": 4_000, "Azure OpenAI|OnDemand": 2_000,
                 "Premium SSD Managed Disks|OnDemand": 9_999}
    # Under actual cost the reserved usage shows as nothing; coverage must not read it.
    actual = {"Dsv5 Series|OnDemand": 3_000, "NDsr H100 v5 Series|SavingsPlan": 0}

    found = finops.analyse(entries_with({"2026-09": month(actual, amortized)}), {}, [], None, None, None, TODAY)

    latest = found["coverage"]["latest"]
    assert (latest["covered"], latest["on_demand"], latest["coverage_pct"]) == (5_000, 5_000, 50.0)
    by_family = {f["family"]: f["coverage_pct"] for f in latest["families"]}
    assert by_family == {"Virtual machines": 25.0, "GPU": 100.0, "Azure OpenAI": 0.0}


def test_a_steady_large_on_demand_meter_is_a_candidate_and_a_spiky_one_is_not():
    months_ = {}
    for index, name in enumerate(CLOSED):
        months_[name] = month({"Esv5 Series|OnDemand": 5_000 + 100 * (index % 2),
                               "Fsv2 Series|OnDemand": 9_000 if index % 2 else 500,
                               "Bs Series|OnDemand": 300})
    months_["2026-10"] = month({"Esv5 Series|OnDemand": 99_999})   # the partial month is never evidence

    found = finops.analyse(entries_with(months_), {}, [], None, None, None, TODAY)

    assert [c["meter"] for c in found["candidates"]] == ["Esv5 Series"]
    candidate = found["candidates"][0]
    assert candidate["monthly_on_demand"] == 5_050 and candidate["annual_on_demand"] == 60_600
    assert list(candidate["months"]) == CLOSED[-6:]


def test_a_meter_this_customer_has_reserved_counts_as_reservable():
    found = finops.analyse(entries_with({
        "2026-08": month({"Azure Databricks Regional|Reservation": 1_000}),
        "2026-09": month({}, {"Azure Databricks Regional|OnDemand": 3_000, "Azure Databricks Regional|Reservation": 1_000}),
    }), {}, [], None, None, None, TODAY)

    assert found["coverage"]["latest"]["families"][0]["family"] == "Other reservable services"


def test_azures_catch_all_meter_is_never_a_candidate_and_database_compute_is_not_a_vm():
    months_ = {name: month({"Unassigned|OnDemand": 50_000, "Unassigned|Reservation": 10,
                            "Flexible Server General Purpose Ddsv5 Series Compute|OnDemand": 4_000,
                            "Flexible Server General Purpose Ddsv5 Series Compute|Reservation": 1})
               for name in CLOSED}

    candidates = finops.analyse(entries_with(months_), {}, [], None, None, None, TODAY)["candidates"]

    assert [(c["meter"], c["family"]) for c in candidates] == [
        ("Flexible Server General Purpose Ddsv5 Series Compute", "Other reservable services")]


def test_a_reservation_bought_up_front_is_found():
    found = finops.analyse(entries_with({
        "2026-08": month({"Dsv5 Series|Reservation": 120_000}, {"Dsv5 Series|Reservation": 10_000}),
        "2026-09": month({"Dsv5 Series|Reservation": 0}, {"Dsv5 Series|Reservation": 10_000}),
    }), {}, [], None, None, None, TODAY)

    assert found["purchases"] == [{"month": "2026-08", "charged": 120_000, "amortized": 10_000,
                                   "bought_up_front": 110_000}]


def test_the_budget_is_the_macc_over_its_term_unless_set():
    macc = {"type": "MACC", "amount": 3_600_000, "currency": "USD", "start": "2026-01-01", "end": "2028-12-31"}
    trend = [{"month": "2026-09", "actual": 120_000, "partial": False}, {"month": "2026-10", "actual": 10, "partial": True}]

    derived = finops.budget_status(finops.monthly_budget(macc, None), trend, 90_000)
    overridden = finops.budget_status(finops.monthly_budget(macc, 150_000), trend, 160_000)

    assert (derived["amount"], derived["source"], derived["term_months"], derived["status"]) == (100_000, "macc", 36, "over")
    assert (overridden["amount"], overridden["source"], overridden["status"]) == (150_000, "set", "forecast-over")
    assert finops.monthly_budget(None, None) is None


def test_growth_compares_the_last_three_closed_months_with_the_three_before():
    merged = {m: {"actual": {"service": {"Storage": 1_000 if m < "2026-07" else 4_000, "VMs": 5_000},
                             "subscription": {}}} for m in CLOSED}

    rows = finops.growth(merged, CLOSED, "service")

    assert rows == [{"name": "Storage", "monthly_now": 4_000, "monthly_before": 1_000, "change": 3_000,
                     "change_pct": 300.0}]


def service_with_customer(fake_api):
    fake_api.get_service = lambda name: None
    store = ConsumptionStore(fake_api)
    book = CustomerBook(store, AccountDirectory(lambda: [dict(a) for a in ACCOUNTS]))
    book.save(None, {"name": "Contoso", "accounts": ["1"],
                     "commitment": {"type": "MACC", "amount": 360_000, "start": "2026-01-01", "end": "2028-12-31"}})
    book.set_scenario("contoso", {"growth": {"Base": 0.02, "Upside": 0.04, "Downside": 0.0}, "active": "Base"})
    store.save_account("1", {"months": {m: month({"Esv5 Series|OnDemand": 5_000}, {"Esv5 Series|OnDemand": 5_000})
                                        for m in CLOSED}, "synced_at": "2026-10-17T00:00:00+00:00"})
    return book, ConsumptionService(store, book, today=lambda: TODAY,
                                    now=lambda: datetime(2026, 10, 17, tzinfo=timezone.utc))


def test_the_customer_view_leaves_out_only_the_forecast_scenarios(fake_api):
    _, service = service_with_customer(fake_api)

    account_view = service.analysis("contoso")
    customer_view = service.analysis("contoso", view="customer")

    assert account_view["scenarios"] and account_view["customer"]["scenario"]
    assert customer_view["scenarios"] is None and customer_view["customer"]["scenario"] is None
    assert customer_view["commitment"] == account_view["commitment"]
    assert customer_view["pipeline"] == account_view["pipeline"]


def test_both_workbooks_carry_finops_and_only_the_account_one_the_scenarios(fake_api):
    _, service = service_with_customer(fake_api)

    account_name, account_book = service.export("contoso")
    customer_name, customer_book = service.export("contoso", view="customer")
    account_sheets = openpyxl.load_workbook(io.BytesIO(account_book)).sheetnames
    customer_sheets = openpyxl.load_workbook(io.BytesIO(customer_book)).sheetnames

    assert "FinOps" in account_sheets and "FinOps" in customer_sheets
    assert "Forecast Scenarios" in account_sheets and "Forecast Scenarios" not in customer_sheets
    assert "for the customer" in customer_name and "for the customer" not in account_name


def test_across_customers_the_largest_steady_on_demand_comes_first(fake_api):
    book, service = service_with_customer(fake_api)
    book.save("contoso", {"name": "Contoso", "accounts": ["1"], "budget_monthly": 4_000})

    rows = service.finops_across()

    assert [(row["name"], row["candidates"], row["candidates_monthly"]) for row in rows] == [("Contoso", 1, 5_000)]
    assert rows[0]["budget"]["source"] == "set" and rows[0]["budget"]["status"] == "over"


def test_the_digest_names_who_is_over_budget_and_the_largest_savings(fake_api):
    book, service = service_with_customer(fake_api)
    book.save("contoso", {"name": "Contoso", "accounts": ["1"], "budget_monthly": 4_000})
    service.prices.discounts = lambda: {"Esv5 Series": {"product": "Virtual Machines Esv5 Series",
                                                        "discounts": {"sp_1y": 0.25, "ri_3y": 0.6}}}

    digest = service.finops_digest()

    assert digest["over_budget"] == 1 and digest["with_savings"] == 1
    assert "Contoso: $3k/mo (1 meters, largest Esv5 Series)" in digest["text"]
    assert "Contoso: $10k in 2026-09 against $4k" in digest["text"]
