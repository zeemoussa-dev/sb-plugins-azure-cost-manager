"""Groups: a parent's MACC carved out to affiliates.

Worth pinning: every affiliate dollar counts against its own MACC and the
parent's; the unallocated rest is the parent's own enrollments' to burn; an
over-allocated parent says so; groups are one level deep; and syncing a parent
brings its affiliates.
"""
from __future__ import annotations

from datetime import date, datetime, timezone

import pytest

from backend.consumption_service import ConsumptionService
from backend.consumption_store import ConsumptionStore
from backend.consumption_sync import ConsumptionSync
from backend.customers import AccountDirectory, CustomerBook, CustomerError

TODAY = date(2026, 10, 17)
ACCOUNTS = [{"name": name, "display_name": display, "agreement_type": "MicrosoftCustomerAgreement",
             "term_start": "2026-01-01", "term_end": "2028-12-31"}
            for name, display in (("1", "Holding"), ("2", "Holding-1-Northwind"), ("3", "Holding-1-Litware"))]


def month(total: float) -> dict:
    spend = {"service": {"VM": total}, "subscription": {"Prod": total}}
    return {"fetched_at": "2026-10-17T00:00:00+00:00", "currency": "USD", "actual": spend, "amortized": spend,
            "meters": {}}


def macc(amount: float) -> dict:
    return {"type": "MACC", "amount": amount, "start": "2026-01-01", "end": "2028-12-31"}


def group(fake_api, northwind_macc=7_000_000, litware_macc=3_000_000):
    fake_api.get_service = lambda name: None
    store = ConsumptionStore(fake_api)
    book = CustomerBook(store, AccountDirectory(lambda: [dict(a) for a in ACCOUNTS]))
    book.save(None, {"name": "Holding", "accounts": ["1"], "commitment": macc(50_000_000)})
    book.save(None, {"name": "Northwind", "accounts": ["2"], "parent": "holding", "commitment": macc(northwind_macc)})
    book.save(None, {"name": "Litware", "accounts": ["3"], "parent": "holding", "commitment": macc(litware_macc)})
    for account, monthly in (("1", 1_000_000), ("2", 500_000), ("3", 250_000)):
        store.save_account(account, {"months": {f"2026-{m:02d}": month(monthly) for m in range(1, 10)},
                                     "synced_at": "2026-10-17T00:00:00+00:00"})
    return book, ConsumptionService(store, book, today=lambda: TODAY)


def test_the_parent_rolls_up_its_affiliates_against_the_whole_macc(fake_api):
    _, service = group(fake_api)

    result = service.analysis("holding")

    assert result["months"][0]["actual"] == 1_750_000
    assert result["commitment"]["consumed"] == 9 * 1_750_000
    assert {account["name"] for account in result["group"]["accounts"]} == {"1", "2", "3"}


def test_each_affiliate_burns_its_own_macc_and_the_rest_is_unallocated(fake_api):
    _, service = group(fake_api)

    allocations = service.analysis("holding")["allocations"]

    assert (allocations["parent_amount"], allocations["allocated"], allocations["unallocated"]) == \
        (50_000_000, 10_000_000, 40_000_000)
    northwind = next(row for row in allocations["affiliates"] if row["name"] == "Northwind")
    assert northwind["burn"]["consumed"] == 4_500_000
    assert northwind["burn"]["amount"] == 7_000_000
    assert allocations["own"] == {"consumed": 9_000_000, "remaining": 31_000_000}
    assert allocations["consumed_total"] == 9 * 1_750_000
    assert allocations["over_allocated"] is False


def test_an_affiliate_sees_its_own_macc_and_the_group_it_belongs_to(fake_api):
    _, service = group(fake_api)

    result = service.analysis("northwind")

    assert result["commitment"]["consumed"] == 4_500_000
    assert result["group"]["parent"] == {"id": "holding", "name": "Holding", "amount": 50_000_000, "currency": "USD",
                                         "entity_customer": None}
    assert result["allocations"] is None


def test_the_parents_export_carries_the_allocation_and_every_enrollment(fake_api):
    import io

    import openpyxl

    _, service = group(fake_api)

    _, data = service.export("holding")
    dashboard = openpyxl.load_workbook(io.BytesIO(data))["Dashboard"]
    values = [cell.value for row in dashboard.iter_rows() for cell in row if cell.value is not None]
    export = openpyxl.load_workbook(io.BytesIO(data))["Export"]
    profiles = {cell.value for cell in export["B"] if cell.value and cell.row > 1}

    assert "MACC Allocation across the Group" in values
    assert "Holding – unallocated (own enrollments)" in values
    assert profiles == {"1", "2", "3"}


def test_affiliates_signing_more_than_the_parent_is_flagged(fake_api):
    _, service = group(fake_api, northwind_macc=45_000_000, litware_macc=10_000_000)

    allocations = service.analysis("holding")["allocations"]

    assert allocations["over_allocated"] is True
    assert allocations["unallocated"] == -5_000_000


def test_groups_are_one_level_deep(fake_api):
    book, _ = group(fake_api)

    with pytest.raises(CustomerError, match="one level"):
        book.save(None, {"name": "Contoso", "accounts": ["1"], "parent": "northwind"})
    book.save(None, {"name": "Contoso", "accounts": ["1"]})
    with pytest.raises(CustomerError, match="affiliates of its own"):
        book.save("holding", {"name": "Holding", "accounts": ["1"], "parent": "contoso"})
    with pytest.raises(CustomerError, match="affiliates"):
        book.delete("holding")


def test_syncing_a_parent_brings_its_affiliates(fake_api):
    book, _ = group(fake_api)
    asked: list[str] = []

    class Azure:
        def query_monthly_cost(self, account, cost_type, first, last):
            asked.append(account)
            return []

        def query_monthly_meters(self, account, first, last, cost_type="ActualCost"):
            return []

        def forecast_monthly(self, account, first, last):
            return []

    store = ConsumptionStore(fake_api)
    sync = ConsumptionSync(store, AccountDirectory(lambda: [dict(a) for a in ACCOUNTS]), lambda: Azure(),
                           today=lambda: TODAY, now=lambda: datetime(2026, 10, 17, tzinfo=timezone.utc),
                           sleep=lambda _: None)
    sync.run("holding")

    assert set(asked) == {"1", "2", "3"}


def _own_group(fake_api):
    book, service = group(fake_api)
    book.save("holding", {"name": "Holding", "accounts": ["1"],
                          "commitment": {**macc(50_000_000), "affiliates": "own"}})
    return book, service


def test_a_parent_whose_affiliates_signed_their_own_is_burned_by_its_own_enrollments_only(fake_api):
    _, service = _own_group(fake_api)

    result = service.analysis("holding")

    # The group is still shown (operator, 2026-10-18: "only the total consumption is affected")...
    assert result["months"][0]["actual"] == 1_750_000
    # ...but only the parent's own enrollments burn its MACC, and nothing is carved out of it.
    assert result["commitment"]["consumed"] == 9 * 1_000_000
    assert result["commitment"]["affiliates"] == "own"
    assert result["allocations"] is None
    assert service.analysis("northwind")["commitment"]["consumed"] == 4_500_000


def test_the_portfolio_burns_an_own_parent_by_its_own_enrollments(fake_api):
    _, service = _own_group(fake_api)

    holding = next(row for row in service.portfolio() if row["id"] == "holding")

    assert holding["burn"]["consumed"] == 9 * 1_000_000


def test_the_report_shows_only_the_enrollments_chosen_and_the_macc_stays_whole(fake_api):
    _, service = group(fake_api)

    hq_only = service.analysis("holding", accounts=["1"])

    assert hq_only["shown_accounts"] == ["1"]
    assert hq_only["months"][0]["actual"] == 1_000_000            # Holding's own enrollment only
    assert hq_only["commitment"]["consumed"] == 9 * 1_750_000     # the MACC still counts the whole group
    assert service.analysis("holding")["shown_accounts"] == ["1", "2", "3"]
    with pytest.raises(ValueError):
        service.analysis("holding", accounts=["999"])


def test_an_affiliates_enrollment_reads_by_its_own_name():
    from backend.customers import short_name

    assert short_name("MDC General Services Holding Company LLC-8990244-Abu Dhabi Future Energy Company - PJSC   Masdar") \
        == "Abu Dhabi Future Energy Company - PJSC Masdar"
    assert short_name("MDC General Services Holding Company LLC") == "MDC General Services Holding Company LLC"
