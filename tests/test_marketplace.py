"""Marketplace: partner solutions per customer and across customers.

Worth pinning: a metered image's hourly lines become one solution; a $0 solution
is adopted, not hidden; an offer marked not eligible stops counting toward the
MACC everywhere -- including the burn-down; a group and its affiliates are not
double-counted across customers; and the itemised charges are checked against
the invoiced Marketplace.
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl

from backend import consumption_analysis as analysis
from backend import marketplace
from backend.consumption_service import ConsumptionService
from backend.consumption_store import ConsumptionStore
from backend.consumption_sync import _aggregate_marketplace
from backend.customers import AccountDirectory, CustomerBook

TODAY = date(2026, 10, 17)


def item(publisher, offer, cost, plan="", subscription="Prod", recurring=False):
    return {"publisher": publisher, "offer": offer, "plan": plan, "subscription": subscription, "cost": cost,
            "quantity": 1.0, "unit": "1 hour", "recurring": recurring, "lines": 1}


def month(items, marketplace_invoiced=None, usage=1000.0):
    invoiced = sum(i["cost"] for i in items) if marketplace_invoiced is None else marketplace_invoiced
    spend = {"service": {"VM": usage}, "subscription": {"Prod": usage}}
    return {"fetched_at": "2026-10-17T00:00:00+00:00", "currency": "USD", "actual": spend, "amortized": spend,
            "meters": {},
            "balance": {"usage": usage, "marketplace": invoiced, "credits": 0.0, "credit_details": [],
                        "fetched_at": "2026-10-17T00:00:00+00:00"},
            "marketplace": {"items": items, "fetched_at": "2026-10-17T00:00:00+00:00"}}


def test_hourly_lines_become_one_solution_per_month():
    charges = [{"publisher": "Northwind", "offer": "Linux", "plan": "8.9", "subscription": "Prod", "cost": 1.5,
                "quantity": 1.0, "unit": "1 hour", "recurring": False}] * 720

    aggregated = _aggregate_marketplace(charges)

    assert len(aggregated) == 1
    assert (aggregated[0]["cost"], aggregated[0]["quantity"], aggregated[0]["lines"]) == (1080.0, 720.0, 720)


def test_a_customer_report_lists_solutions_adopted_ones_and_what_counts():
    entries = {"1": {"months": {
        "2026-08": month([item("Northwind", "Linux", 300, "8.9"), item("Litware", "Free plan", 0)]),
        "2026-09": month([item("Northwind", "Linux", 200, "8.9"), item("Contoso", "Backup SaaS", 100, recurring=True)]),
    }}}

    report = marketplace.customer_report(entries, {"1": "Enrollment"}, {"Contoso|Backup SaaS"}, "2026-01", "2026-10")

    assert [s["offer"] for s in report["solutions"]] == ["Linux", "Backup SaaS"]
    assert [s["offer"] for s in report["adopted"]] == ["Free plan"]
    assert (report["total"], report["macc"], report["not_macc"]) == (600, 500, 100)
    assert report["solutions"][1]["kind"] == "SaaS" and report["solutions"][1]["macc_eligible"] is False
    assert report["publishers"][0] == {"publisher": "Northwind", "total": 500, "macc": 500, "solutions": 1}
    assert report["reconciliation"] == []


def test_from_the_start_names_the_first_month_on_record():
    entries = {"1": {"months": {"2024-03": month([item("Northwind", "Linux", 10)]),
                                "2026-09": month([item("Northwind", "Linux", 20)])}}}

    report = marketplace.customer_report(entries, {}, set(), None, "2026-10")

    assert report["period"] == {"first": "2024-03", "last": "2026-10", "from_start": True}
    assert report["total"] == 30


def test_a_month_whose_charges_disagree_with_the_invoice_is_reported():
    entries = {"1": {"months": {"2026-09": month([item("Northwind", "Linux", 200)], marketplace_invoiced=350)}}}

    gaps = marketplace.reconcile(entries, None, None)

    assert gaps == [{"account": "1", "month": "2026-09", "itemised": 200, "invoiced": 350}]


def test_an_offer_marked_not_eligible_stops_burning_the_macc():
    entries = {"1": {"months": {"2026-09": month([item("Northwind", "Linux", 200), item("Contoso", "Backup", 100)])}}}
    merged, _ = analysis.merge(entries, None, {"Contoso|Backup"})
    commitment = {"type": "MACC", "amount": 10_000, "currency": "USD", "start": "2026-01-01", "end": "2026-12-31",
                  "include_marketplace": True, "subtract_credits": True}

    result = analysis.analyse(merged, "USD", commitment, None, TODAY)

    assert result["months"][0]["actual"] == 1300          # the dashboard still shows all Marketplace
    assert result["commitment"]["consumed"] == 1200       # the MACC does not burn the excepted offer


def test_across_customers_a_group_is_not_double_counted():
    shared = {"2": {"months": {"2026-09": month([item("Northwind", "Linux", 100)])}}}
    parent_own = {"1": {"months": {"2026-09": month([item("Northwind", "Linux", 50), item("Litware", "AI", 25)])}}}
    customers = [({"id": "affiliate", "name": "Affiliate", "parent": "group"}, shared),
                 ({"id": "group", "name": "Group", "parent": None}, {**parent_own, **shared})]

    report = marketplace.across_customers(customers, set(), "2026-01", "2026-12")

    assert report["total"] == 175
    linux = next(s for s in report["solutions"] if s["offer"] == "Linux")
    assert linux["customers"] == [{"name": "Affiliate", "total": 100}, {"name": "Group", "total": 50}]
    assert linux["adopted_by"] == 2


def test_the_marketplace_reaches_both_exports(fake_api):
    fake_api.get_service = lambda name: None
    store = ConsumptionStore(fake_api)
    book = CustomerBook(store, AccountDirectory(lambda: [{"name": "1", "display_name": "Contoso",
                                                          "agreement_type": "EnterpriseAgreement",
                                                          "term_start": "2025-01-01", "term_end": "2027-12-31"}]))
    book.save(None, {"name": "Contoso", "accounts": ["1"]})
    store.save_account("1", {"months": {f"2026-{m:02d}": month([item("Northwind", "Linux", 100 * m, "8.9")])
                                        for m in range(1, 10)}, "synced_at": "2026-10-17T00:00:00+00:00"})
    service = ConsumptionService(store, book, today=lambda: TODAY)
    service.set_macc_eligible("Northwind", "Linux", False)

    _, customer_workbook = service.export("contoso")
    _, across_workbook = service.export_marketplace()
    sheet = openpyxl.load_workbook(io.BytesIO(customer_workbook))["Marketplace"]
    across = openpyxl.load_workbook(io.BytesIO(across_workbook))

    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert "Linux" in values and "no" in values
    assert across.sheetnames == ["Marketplace Summary", "Solutions", "Solution x Customer"]
    assert service.marketplace_across()["not_macc"] == 4500
    # The customer workbook's Marketplace sheet runs from the first month on record.
    assert sheet["A2"].value.startswith("From 2026-01 to 2026-10")


def test_exports_carry_the_operators_theme_so_charts_take_its_colours(fake_api):
    import zipfile

    from backend import export_xlsx

    data = export_xlsx.build_marketplace({"period": {"first": "2026-01", "last": "2026-09", "from_start": False},
                                          "total": 0, "macc": 0, "not_macc": 0, "solutions": [], "publishers": [],
                                          "customers": [], "monthly": []}, TODAY)
    theme = zipfile.ZipFile(io.BytesIO(data)).read("xl/theme/theme1.xml").decode("utf8")

    assert "156082" in theme and "E97132" in theme   # accent1 and accent2 of the operator's workbook
    assert "Aptos" in theme and "4F81BD" not in theme  # not openpyxl's 2007 default
