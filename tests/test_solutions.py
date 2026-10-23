"""What each subscription is used for.

Worth pinning: a subscription is placed by the services it runs and its names,
residue does not count, plain VMs show only when nothing else fits, a pitch is
the piece missing from the whole estate, the catalog can be replaced but not
broken, and the sync reads the footprint once per closed month.
"""
from __future__ import annotations

from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook
import io

from backend import export_reports, solutions
from backend.consumption_store import ConsumptionStore
from backend.consumption_sync import ConsumptionSync
from backend.customers import AccountDirectory, CustomerBook

CATALOG = solutions.patterns(None)


def entry(subscriptions: dict[str, dict]) -> dict:
    """Costs are three months' totals, as the sync keeps them."""
    return {"footprint": {"first_month": "2026-07", "last_month": "2026-09", "months": 3, "fetched_at": "x",
                          "subscriptions": {name: {"services": found.get("services", {}),
                                                   "meters": found.get("meters", {}),
                                                   "resource_groups": found.get("groups", {})}
                                            for name, found in subscriptions.items()}}}


def placed(view: dict, subscription: str) -> list[str]:
    return [found["id"] for row in view["subscriptions"] if row["subscription"] == subscription
            for found in row["solutions"]]


def test_a_subscription_is_placed_by_its_services_and_names():
    view = solutions.customer_view({"1": entry({
        "sub-data-prod": {"services": {"Microsoft Fabric": 30000, "Azure Data Factory v2": 6000, "Storage": 9000},
                          "meters": {"Fabric Capacity": 30000}, "groups": {"rg-datalake-prod": 20000}},
        "sub-ai": {"services": {"Foundry Models": 9000, "Azure App Service": 3000},
                   "meters": {"Azure OpenAI": 9000}},
        "sub-vms": {"services": {"Virtual Machines": 12000, "Storage": 3000}},
        "sub-tiny": {"services": {"Azure Databricks": 30}},
    })}, CATALOG)

    lakehouse = next(row for row in view["subscriptions"] if row["subscription"] == "sub-data-prod")
    assert placed(view, "sub-data-prod") == ["lakehouse"]           # pipelines are part of it, not a second solution
    assert lakehouse["solutions"][0]["confidence"] == solutions.STRONG and lakehouse["environment"] == "Production"
    assert placed(view, "sub-ai") == ["genai"]
    assert placed(view, "sub-vms") == ["iaas"]                      # nothing more specific fits
    assert "sub-tiny" not in [row["subscription"] for row in view["subscriptions"]]   # $10 a month is residue


def test_a_pitch_is_the_piece_missing_from_the_whole_estate():
    view = solutions.customer_view({"1": entry({
        "sub-ai": {"services": {"Foundry Models": 9000}, "meters": {"Azure OpenAI": 9000}},
        "sub-search": {"services": {"Azure Cognitive Search": 3000}},
    })}, CATALOG)

    pitches = {pitch["pitch"] for pitch in view["pitches"] if pitch["solution"] == "Generative AI application"}
    assert not any("AI Search" in pitch for pitch in pitches)        # they run it, in another subscription
    assert any("API Management" in pitch for pitch in pitches) and any("Compass" in pitch for pitch in pitches)


def test_across_customers_counts_each_solution_once_per_customer():
    view = solutions.customer_view({"1": entry({"a": {"services": {"Sentinel": 9000}},
                                                "b": {"services": {"Sentinel": 3000}}})}, CATALOG)
    report = solutions.across([({"id": "c", "name": "Contoso"}, view), ({"id": "n", "name": "Never synced"},
                                                                       solutions.customer_view({}, CATALOG))])

    soc = next(item for item in report["by_solution"] if item["id"] == "soc")
    assert (soc["customers"], soc["subscriptions"], soc["monthly"]) == (["Contoso"], 2, 4000.0)
    assert [row["name"] for row in report["customers"]] == ["Contoso"]
    workbook = load_workbook(io.BytesIO(export_reports.build_solutions(report, date(2026, 10, 23))))
    assert workbook.sheetnames == ["Solutions", "Subscriptions", "Pitches", "Not Placed"]


def test_the_catalog_can_be_replaced_but_not_broken():
    own = [{"id": "sap", "name": "SAP", "play": solutions.INFRA, "requires": [["name:sap"]]}]
    assert [pattern["id"] for pattern in solutions.patterns(own)] == ["sap"]
    for broken in ([], [{"id": "x", "name": "X", "play": "Nope", "requires": [["service:A"]]}],
                   [{"id": "x", "name": "X", "play": solutions.AI, "requires": [["vm:A"]]}],
                   [{"id": "x", "name": "X", "play": solutions.AI, "requires": []}]):
        with pytest.raises(solutions.PatternError):
            solutions.validated(broken)


def test_the_sync_reads_the_footprint_once_per_closed_month(fake_api):
    asked: list[tuple] = []

    class Azure:
        def list_tag_keys(self, account):
            return []

        def query_by_subscription(self, account, first, last, dimension):
            asked.append((dimension, first, last))
            return [{"month": "2026-09", "subscription": "sub-ai", "name": "Foundry Models", "cost": 900.0},
                    {"month": "2026-08", "subscription": "sub-ai", "name": "Foundry Models", "cost": 600.0}]

    fake_api.get_service = lambda name: None
    store = ConsumptionStore(fake_api)
    directory = AccountDirectory(lambda: [{"name": "1", "display_name": "Contoso", "agreement_type": "EnterpriseAgreement",
                                           "term_start": "2026-01-01", "term_end": "2028-12-31"}])
    CustomerBook(store, directory).save(None, {"name": "Contoso", "accounts": ["1"]})

    def sync_on(day: date) -> ConsumptionSync:
        return ConsumptionSync(store, directory, lambda: Azure(), today=lambda: day,
                               now=lambda: datetime(day.year, day.month, day.day, tzinfo=timezone.utc),
                               sleep=lambda _: None)

    sync_on(date(2026, 10, 17)).run_tags()
    sync_on(date(2026, 10, 30)).run_tags()         # same closed month: nothing asked
    first = list(asked)
    sync_on(date(2026, 11, 2)).run_tags()          # October closed: asked again

    assert sorted(first) == [("MeterSubCategory", "2026-07", "2026-09"), ("ResourceGroupName", "2026-07", "2026-09"),
                             ("ServiceName", "2026-07", "2026-09")]
    assert sorted(asked[len(first):])[0] == ("MeterSubCategory", "2026-08", "2026-10")
    footprint = store.account("1")["footprint"]
    assert footprint["subscriptions"]["sub-ai"]["services"] == {"Foundry Models": 1500.0}
