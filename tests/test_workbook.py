"""The operator's workbook, per customer: measures, plan lines, forecasts, the
pipeline and the export.

Worth pinning: Marketplace and credits are switches (and the signed value counts
them its own way), a GPU or PTU line is recognised from its meter, a forecast
always carries how far to trust it, the pipeline adds only what is still
expected, and the export is a workbook Excel can open.
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl

from backend import categories, consumption_analysis as analysis, forecasting
from backend.consumption_service import ConsumptionService
from backend.consumption_store import ConsumptionStore
from backend.customers import AccountDirectory, CustomerBook
from backend.pipeline import Pipeline, by_month

TODAY = date(2026, 10, 16)


def month(total: float, balance: dict | None = None, meters: dict | None = None) -> dict:
    spend = {"service": {"VM": total}, "subscription": {"Prod": total}}
    data = {"fetched_at": "2026-10-16T00:00:00+00:00", "currency": "USD", "actual": spend, "amortized": spend,
            "meters": meters or {}}
    if balance:
        data["balance"] = {**balance, "fetched_at": "2026-10-16T00:00:00+00:00"}
    return data


def balance(usage, marketplace=0.0, credits=0.0):
    return {"usage": usage, "marketplace": marketplace, "credits": credits,
            "credit_details": [{"name": "Contract Price Match Credit", "value": credits}] if credits else [],
            "currency": "USD"}


def test_marketplace_and_credits_are_switches():
    merged, _ = analysis.merge({"ea": {"months": {"2026-06": month(3_678_174, balance(3_677_821, 353, 288_493))}}})

    def total(marketplace, credits):
        return analysis.monthly_totals(merged, "2026-10", {"marketplace": marketplace, "credits": credits})[0]["actual"]

    assert total(True, True) == 3_389_681       # the partner export's Total
    assert total(False, True) == 3_389_328      # what counts toward quota
    assert total(True, False) == 3_678_174
    assert total(False, False) == 3_677_821


def test_an_enrollment_without_a_balance_is_measured_by_cost_management():
    merged, _ = analysis.merge({"mca": {"months": {"2026-06": month(500)}}})

    assert analysis.monthly_totals(merged, "2026-10", {"marketplace": False, "credits": True})[0]["actual"] == 500


def test_the_signed_value_counts_marketplace_its_own_way():
    merged, _ = analysis.merge({"ea": {"months": {"2026-06": month(1100, balance(1000, 100, 0))}}})
    commitment = {"type": "MACC", "amount": 10_000, "currency": "USD", "start": "2026-06-01", "end": "2027-05-31",
                  "include_marketplace": True, "subtract_credits": True}

    result = analysis.analyse(merged, "USD", commitment, None, TODAY, measure={"marketplace": False, "credits": True})

    assert result["months"][0]["actual"] == 1000         # the dashboard, without Marketplace
    assert result["commitment"]["consumed"] == 1100      # the MACC still burns it


def test_gpu_and_openai_lines_are_recognised_from_their_meters():
    assert categories.classify("NDsr H100 v5 Series", "Reservation") == "GPU H100"
    assert categories.classify("NC A100 v4 Series", "OnDemand") == "GPU A100"
    assert categories.classify("NVads A10 v5 Series", "OnDemand") == "GPU A10"
    assert categories.classify("Azure OpenAI", "Reservation") == "Azure OpenAI PTU"
    assert categories.classify("Azure OpenAI", "OnDemand") == "Azure OpenAI pay-as-you-go"
    assert categories.classify("Dsv5 Series", "OnDemand") is None

    mix = categories.mix(1000, {"NDsr H100 v5 Series|Reservation": 200, "Azure OpenAI|Reservation": 100}, 50)

    assert mix == {"GPU H100": 200, "Azure OpenAI PTU": 100, "Marketplace": 50, "Azure (other)": 650}
    assert sum(mix.values()) == 1000


def trend_of(values: dict[str, float]) -> list[dict]:
    return [{"month": m, "actual": v, "partial": m >= "2026-10", "mom_pct": None} for m, v in values.items()]


def test_the_automatic_forecast_follows_the_trend_and_says_how_far_to_trust_it():
    history = {f"{2025 + (i // 12)}-{i % 12 + 1:02d}": 1000 + 100 * i for i in range(21)}  # Jan 2025 .. Sep 2026
    result = forecasting.automatic(trend_of(history), TODAY)

    assert result["r_squared"] > 0.99
    assert result["mape"] < 0.01
    assert result["slope"] == 100.0
    october = next(row for row in result["series"] if row["month"] == "2026-10")
    assert round(october["forecast"]) == 3100
    assert result["next_year"] == 2027


def test_no_automatic_forecast_from_too_little_history():
    assert forecasting.automatic(trend_of({"2026-08": 10, "2026-09": 12}), TODAY) is None


def test_scenarios_grow_from_the_last_closed_month_and_add_an_uplift():
    trend = trend_of({"2026-08": 900, "2026-09": 1000})
    result = forecasting.scenarios(trend, TODAY, {"seasonality": False, "uplift_usd": 500, "uplift_from": "2026-12",
                                                  "growth": {"Base": 0.1, "Upside": 0.2, "Downside": 0.0}})

    october, december = result["rows"][0], result["rows"][2]
    assert (october["month"], october["Base"], october["Downside"]) == ("2026-10", 1100.0, 1000.0)
    assert december["Base"] == round(1000 * 1.1 ** 3 + 500, 2)


def test_the_pipeline_adds_only_what_is_still_expected():
    opportunities = [
        {"status": "Open", "plan": [{"month": "2026-11", "usd": 50_000}, {"month": "2026-12", "usd": 100_000}]},
        {"status": "Won", "plan": [{"month": "2026-12", "usd": 10_000}]},
        {"status": "Lost", "plan": [{"month": "2026-12", "usd": 999_999}]},
    ]

    assert by_month(opportunities) == {"2026-11": 50_000, "2026-12": 110_000}


class FakeOpportunities:
    def __init__(self):
        self.created = []

    def customers(self):
        return [{"name": "Contoso", "stem": "Contoso", "tag": "customer/contoso", "affiliates": []}]

    def list_for_customer(self, stem, include_affiliates=True):
        # The shape the Entities plugin returns: partner and product are lists.
        return [{"id": "o1", "title": "Data platform", "status": "Open", "customer": "Contoso", "affiliate": None,
                 "partner": ["Northwind", "Litware"], "product": ["azure"],
                 "plan": [{"month": "2026-11", "usd": 25_000}], "expected_monthly_usd": 25_000}]

    def create(self, owner_stem, title, fields, plan):
        self.created.append((owner_stem, title, fields, plan))
        return {"id": "o2", "title": title}

    def update(self, opportunity_id, fields, plan=None):
        return {"id": opportunity_id, **fields}


def built(fake_api, with_entities=True):
    service = FakeOpportunities()
    fake_api.get_service = lambda name: service if with_entities and name == "entities.opportunities" else None
    store = ConsumptionStore(fake_api)
    directory = AccountDirectory(lambda: [{"name": "1", "display_name": "Contoso", "agreement_type": "EnterpriseAgreement",
                                          "term_start": "2025-01-01", "term_end": "2027-12-31"}])
    book = CustomerBook(store, directory)
    book.save(None, {"name": "Contoso", "accounts": ["1"], "entity_customer": "Contoso",
                     "commitment": {"type": "MACC", "amount": 30_000, "start": "2025-01-01", "end": "2027-12-31"}})
    history = {f"{2025 + (i // 12)}-{i % 12 + 1:02d}": month(800 + 20 * i, balance(800 + 20 * i, 5, 0),
                                                           {"NDsr H100 v5 Series|Reservation": 100})
               for i in range(22)}
    store.save_account("1", {"months": history, "synced_at": "2026-10-16T00:00:00+00:00"})
    return ConsumptionService(store, book, today=lambda: TODAY, pipeline=Pipeline(fake_api)), service


def test_the_pipeline_comes_from_the_entities_plugin_and_feeds_the_burn_down(fake_api):
    consumption, _ = built(fake_api)

    result = consumption.analysis("contoso")

    assert result["pipeline"]["opportunities"][0]["title"] == "Data platform"
    assert result["commitment"]["pipeline_in_term"] == 25_000
    assert result["commitment"]["projected_with_pipeline"] == result["commitment"]["projected_at_end"] + 25_000


def test_without_the_entities_plugin_the_pipeline_says_why_and_the_rest_works(fake_api):
    consumption, _ = built(fake_api, with_entities=False)

    result = consumption.analysis("contoso")

    assert result["pipeline"]["available"] is False
    assert "Entities" in result["pipeline"]["reason"]
    assert result["dashboard"]["kpis"]["all_time"] > 0


def test_the_export_is_a_workbook_with_the_dashboard_and_its_data(fake_api):
    consumption, _ = built(fake_api)

    name, data = consumption.export("contoso")
    workbook = openpyxl.load_workbook(io.BytesIO(data))

    assert name.endswith(".xlsx") and "Contoso" in name
    # The operator's workbook first, in its order; the plugin's extra sheets after.
    assert workbook.sheetnames[:5] == ["Dashboard", "Forecast Auto", "Forecast Scenarios", "Pipeline", "Export"]
    assert {"Services", "Subscriptions", "GPU & AI", "Commitment"} <= set(workbook.sheetnames)

    dashboard = workbook["Dashboard"]
    assert len(dashboard._charts) == 10
    assert dashboard.sheet_view.showGridLines is False
    assert dashboard["A3"].value == "TOTAL CONSUMPTION – ALL TIME"
    assert dashboard["A4"].fill.fgColor.rgb.endswith("EEF3FB") and dashboard["A4"].font.name == "Aptos Narrow"
    assert dashboard["A10"].fill.fgColor.rgb.endswith("FDF2E9")  # the signed value's peach tiles
    assert dashboard["L3"].value == 30_000

    pipeline = workbook["Pipeline"]
    assert (pipeline["B3"].value, pipeline["C3"].value, pipeline["E3"].value) == ("Data platform", "Northwind, Litware", "Open")
    assert pipeline["F1"].value == 2026 and pipeline["P2"].value == "Nov" and pipeline["P3"].value == 25_000

    export = workbook["Export"]
    assert export["D2"].value == "Total" and export["C2"].value == 2026
    assert export["D3"].value == "January"
    # January 2026: usage 1040, of which 100 on H100 -- so Azure 940 + GPU 100 + Marketplace 5 = 1045.
    assert (export["G3"].value, export["H3"].value, export["E3"].value, export["L3"].value) == (940, 100, 5, 1045)
